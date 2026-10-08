from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
DOC_PATH = ROOT / "เอกสารประกอบการพัฒนาซอฟต์แวร์_ฉบับเพิ่มภาคผนวก.docx"
ASSET_DIR = ROOT / "appendix_assets"


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def style_run(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)


def add_text(document, text, *, size=16, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             before=0, after=4, keep_with_next=False):
    paragraph = document.add_paragraph()
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.15
    paragraph.paragraph_format.keep_with_next = keep_with_next
    style_run(paragraph.add_run(text), size=size, bold=bold)
    return paragraph


def add_step(document, number, text):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.first_line_indent = Inches(-0.22)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.1
    style_run(paragraph.add_run(f"{number}. "), size=15, bold=True)
    style_run(paragraph.add_run(text), size=15)


def add_figure(document, image_name, caption):
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_before = Pt(6)
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(str(ASSET_DIR / image_name), width=Inches(6.15))
    caption_paragraph = document.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_after = Pt(4)
    style_run(caption_paragraph.add_run(caption), size=14)


def add_section(document, title, description, steps, image_name, caption, page_break=True):
    if page_break:
        document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    add_text(document, title, size=18, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT,
             after=4, keep_with_next=True)
    add_text(document, description, size=16, after=4)
    for index, step in enumerate(steps, 1):
        add_step(document, index, step)
    add_figure(document, image_name, caption)


document = Document(DOC_PATH)

# Replace the appendix placeholders while preserving the rest of the report.
for paragraph in list(document.paragraphs):
    text = paragraph.text.strip()
    if text.startswith("(พื้นที่สำหรับจัดทำคู่มือ") or text.startswith("(เว้นว่างไว้สำหรับรูปภาพ"):
        remove_paragraph(paragraph)

add_text(
    document,
    "ภาคผนวกนี้แสดงขั้นตอนการใช้งานระบบ BRU Fondue สำหรับผู้แจ้งปัญหา "
    "ผู้ดูแลระบบ และช่างผู้รับผิดชอบ โดยใช้ภาพจากระบบที่ติดตั้งใช้งานจริงบน Render "
    "และเชื่อมต่อฐานข้อมูล Neon PostgreSQL",
    size=16,
    after=8,
)

add_section(
    document,
    "1 การเข้าสู่ระบบ",
    "ผู้ใช้ทุกบทบาทเข้าสู่ระบบด้วยอีเมลที่ได้รับการลงทะเบียนและรหัสผ่านของตนเอง",
    [
        "เปิดเว็บไซต์ BRU Fondue แล้วเลือกเมนู เข้าสู่ระบบ",
        "กรอกอีเมลและรหัสผ่านให้ครบถ้วน",
        "กดปุ่ม เข้าสู่ระบบ ระบบจะนำผู้ใช้ไปยังพื้นที่ทำงานตามสิทธิ์ของบัญชี",
    ],
    "01-login.png",
    "ภาพที่ 1 หน้าจอเข้าสู่ระบบ BRU Fondue",
    page_break=False,
)

add_section(
    document,
    "2 การสร้างคำร้องแจ้งปัญหาสำหรับผู้แจ้ง",
    "ผู้แจ้งสามารถบันทึกประเภทปัญหา รายละเอียด สถานที่ พิกัด และรูปภาพประกอบผ่านแบบฟอร์มเดียว",
    [
        "เลือกหมวดหมู่ปัญหาและระบุหัวข้อพร้อมรายละเอียดที่พบ",
        "เลือกอาคารหรือห้อง และตรวจสอบตำแหน่งบนแผนที่",
        "แนบรูปภาพ เลือกระดับความเร่งด่วน แล้วตรวจสอบข้อมูลก่อนส่งคำร้อง",
    ],
    "02-reporter-create.png",
    "ภาพที่ 2 แบบฟอร์มสร้างคำร้องแจ้งปัญหาสำหรับ Reporter",
)

add_section(
    document,
    "3 การตรวจสอบภาพรวมสำหรับผู้ดูแลระบบ",
    "หน้า Dashboard ช่วยให้ผู้ดูแลตรวจสอบจำนวนคำร้อง สถานะงาน ระยะเวลาดำเนินการ "
    "และผลการให้บริการได้จากจุดเดียว",
    [
        "กำหนดช่วงเวลา สถานะ อาคาร หรือหมวดหมู่ที่ต้องการตรวจสอบ",
        "อ่านค่าจากการ์ดสรุปและรายการสถิติด้านล่าง",
        "ดาวน์โหลดข้อมูลเป็น CSV หรือพิมพ์รายงาน PDF เมื่อต้องการนำไปวิเคราะห์ต่อ",
    ],
    "03-admin-dashboard.png",
    "ภาพที่ 3 หน้าจอสถิติและรายงานงานซ่อมสำหรับ Admin",
)

add_section(
    document,
    "4 การจัดการและมอบหมายคำร้องสำหรับผู้ดูแลระบบ",
    "ผู้ดูแลใช้หน้าจัดการคำร้องเพื่อตรวจสอบตำแหน่ง คัดกรองรายการ รับคำร้อง "
    "และมอบหมายงานให้ช่างที่เหมาะสม",
    [
        "ตรวจสอบตำแหน่งคำร้องจากแผนที่และใช้ตัวกรองเพื่อค้นหารายการ",
        "เปิดรายละเอียดคำร้อง ตรวจสอบข้อมูลผู้แจ้ง รูปภาพ และความเร่งด่วน",
        "รับคำร้องและเลือกช่างผู้รับผิดชอบก่อนยืนยันการมอบหมายงาน",
    ],
    "04-admin-tickets.png",
    "ภาพที่ 4 ศูนย์ควบคุมและจ่ายงานซ่อมบำรุงสำหรับ Admin",
)

add_section(
    document,
    "5 การรับงานและอัปเดตสถานะสำหรับช่าง",
    "ช่างตรวจสอบงานที่ได้รับมอบหมาย เปิดรายละเอียดหน้างาน และรายงานความคืบหน้าจนงานเสร็จสิ้น",
    [
        "เปิดเมนู งานของฉัน เพื่อตรวจสอบรายการที่ได้รับมอบหมาย",
        "เลือกงานเพื่อดูสถานที่ รายละเอียดปัญหา รูปภาพ และข้อมูลผู้เกี่ยวข้อง",
        "ปรับสถานะตามขั้นตอนการทำงาน แนบภาพหลังซ่อม และบันทึกผลเมื่อดำเนินการเสร็จ",
    ],
    "05-technician-jobs.png",
    "ภาพที่ 5 หน้าจองานที่ได้รับมอบหมายสำหรับ Technician",
)

document.save(DOC_PATH)
print(DOC_PATH)
