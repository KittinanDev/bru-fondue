from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = Path(r"C:\Users\Kittinan\Downloads\asd\เอกสารประกอบการพัฒนาซอฟต์แวร์ - รายวิชาวิศวกรรมซอฟต์แวร์.docx")
OUTPUT = ROOT / "เอกสารประกอบการพัฒนาซอฟต์แวร์_คู่มือฉบับสมบูรณ์.docx"
ASSETS = ROOT / "manual_assets"

def font(run, size=16, bold=False):
    run.font.name = "TH Sarabun New"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "TH Sarabun New")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)

def paragraph(doc, text="", size=16, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    font(p.add_run(text), size, bold)
    return p

def heading(doc, text, level=1):
    sizes = {1: 20, 2: 18, 3: 16}
    return paragraph(doc, text, sizes[level], True, WD_ALIGN_PARAGRAPH.LEFT, 6, 5, True)

def step(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.28)
    p.paragraph_format.first_line_indent = Inches(-.28)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.08
    font(p.add_run(f"{number}. "), 15, True)
    font(p.add_run(text), 15)

def figure(doc, filename, caption, width=5.95):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    c = paragraph(doc, caption, 14, False, WD_ALIGN_PARAGRAPH.CENTER, 0, 5)
    c.paragraph_format.keep_with_next = False

def page(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def section(doc, title, intro, steps, image, caption, new_page=True, width=5.95):
    if new_page: page(doc)
    heading(doc, title, 2)
    paragraph(doc, intro)
    for i, item in enumerate(steps, 1): step(doc, i, item)
    figure(doc, image, caption, width)

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell(cell, text, bold=False, white=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    font(r, 14, bold)
    if white: r.font.color.rgb = RGBColor(255, 255, 255)

doc = Document(SOURCE)

# Remove the unfinished appendix block from the source while keeping Chapters 1-5 intact.
body = doc._element.body
start = None
for p in doc.paragraphs:
    if p.text.strip() == "ภาคผนวก":
        start = p._element
        break
if start is not None:
    children = list(body)
    index = children.index(start)
    for element in children[index:]:
        if element.tag != qn("w:sectPr"):
            body.remove(element)

# Remove empty trailing paragraphs/page breaks left before the old appendix.
while True:
    content = [element for element in list(body) if element.tag != qn("w:sectPr")]
    if not content or content[-1].tag != qn("w:p"):
        break
    texts = content[-1].xpath(".//w:t")
    if any((node.text or "").strip() for node in texts):
        break
    body.remove(content[-1])

page(doc)
paragraph(doc, "ภาคผนวก", 22, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 10, True)
paragraph(doc, "คู่มือการใช้งานระบบ BRU Fondue", 20, True, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, True)
paragraph(doc, "คู่มือนี้อธิบายการใช้งานระบบรับแจ้งและติดตามงานซ่อมของมหาวิทยาลัยราชภัฏบุรีรัมย์ ครอบคลุมผู้แจ้งปัญหา เจ้าหน้าที่ผู้ดูแลระบบ และช่างซ่อมบำรุง ภาพประกอบบันทึกจากระบบที่เผยแพร่บน Render และเชื่อมต่อฐานข้อมูล Neon PostgreSQL", 16)
heading(doc, "ขอบเขตผู้ใช้งาน", 2)
for label, text in [
    ("ผู้แจ้งปัญหา", "สร้างคำร้อง ระบุสถานที่ แนบภาพ ติดตามสถานะ รับการแจ้งเตือน และประเมินงานซ่อม"),
    ("ผู้ดูแลระบบ", "ตรวจสอบคำร้อง รับเรื่อง มอบหมายหรือเปลี่ยนช่าง ปฏิเสธหรือยกเลิกคำร้อง ดูรายงาน และจัดการข้อมูลพื้นฐาน"),
    ("ช่างซ่อมบำรุง", "เปิดดูงานที่ได้รับมอบหมาย เปลี่ยนสถานะ บันทึกผล และแนบภาพหลังซ่อม"),
]:
    p = doc.add_paragraph(style=None); p.paragraph_format.space_after = Pt(3)
    font(p.add_run(f"• {label}: "), 15, True); font(p.add_run(text), 15)
heading(doc, "ข้อกำหนดเบื้องต้น", 2)
for i, text in enumerate(["ใช้อุปกรณ์ที่เชื่อมต่ออินเทอร์เน็ตและเว็บเบราว์เซอร์รุ่นปัจจุบัน", "มีบัญชีที่ผู้ดูแลระบบสร้างไว้และทราบอีเมลกับรหัสผ่านของตน", "อนุญาตการเข้าถึงตำแหน่งเมื่อจำเป็นต้องปักหมุด และเตรียมรูปภาพ JPG PNG หรือ WebP ขนาดไม่เกิน 5 MB"], 1): step(doc, i, text)

section(doc, "1 การเปิดเว็บไซต์และทำความรู้จักหน้าหลัก", "หน้าหลักอธิบายวัตถุประสงค์ ประโยชน์ และขั้นตอนหลักของระบบ ผู้ใช้เริ่มแจ้งปัญหาหรือเปิดส่วนวิธีใช้งานได้จากหน้านี้", ["เปิดเว็บไซต์ BRU Fondue", "เลือก แจ้งปัญหาใหม่ เพื่อเริ่มสร้างคำร้อง หรือเลือก ดูวิธีการใช้งาน", "ผู้ใช้ที่มีบัญชีเลือก เข้าสู่ระบบ จากมุมขวาบน"], "01-landing.png", "ภาพที่ 1 หน้าหลัก BRU Fondue")
section(doc, "2 การเข้าสู่ระบบ", "ผู้ใช้ทุกบทบาทเข้าสู่ระบบด้วยบัญชีที่ได้รับจากผู้ดูแล ระบบจะนำไปยังพื้นที่ทำงานตามสิทธิ์โดยอัตโนมัติ", ["เปิดเมนู เข้าสู่ระบบ", "กรอกอีเมลและรหัสผ่าน", "เลือกแสดงรหัสผ่านเพื่อตรวจสอบได้หากจำเป็น", "กด เข้าสู่ระบบ หากข้อมูลไม่ถูกต้องให้ตรวจอีเมลและติดต่อผู้ดูแลเพื่อรีเซ็ตรหัสผ่าน"], "02-login.png", "ภาพที่ 2 หน้าจอเข้าสู่ระบบ")
section(doc, "3 การสร้างคำร้องแจ้งปัญหา", "แบบฟอร์มแบ่งเป็น 3 ขั้นตอนเพื่อลดความผิดพลาด ผู้แจ้งสามารถย้อนกลับไปแก้ข้อมูลก่อนส่งได้", ["ขั้นอธิบายปัญหา เลือกหมวดหมู่ กรอกหัวข้อ รายละเอียด และระดับความเร่งด่วน", "ขั้นระบุสถานที่ กรอกอาคาร ชั้น ห้อง หรือบริเวณ และตรวจสอบหมุดบนแผนที่", "ขั้นรูปภาพ แนบภาพก่อนซ่อมที่เห็นปัญหาชัดเจน", "ตรวจข้อมูลแล้วกดส่งคำร้อง ระบบจะแสดงรหัสคำร้องสำหรับติดตาม"], "03-reporter-create.png", "ภาพที่ 3 แบบฟอร์มสร้างคำร้องสำหรับผู้แจ้ง")
section(doc, "4 การติดตามคำร้องของผู้แจ้ง", "หน้า คำร้องของฉัน แสดงรายการทั้งหมดของบัญชีปัจจุบัน พร้อมรหัส หัวข้อ สถานที่ และสถานะล่าสุด", ["เปิดเมนู คำร้อง หรือ คำร้องของฉัน", "ใช้สถานะและข้อมูลบนการ์ดเพื่อดูภาพรวม", "เลือกคำร้องที่ต้องการเพื่อเปิดรายละเอียด", "คำร้องที่ยังรอรับเรื่องสามารถยกเลิกได้จากหน้ารายละเอียด"], "04-reporter-tickets.png", "ภาพที่ 4 รายการคำร้องของผู้แจ้ง")
section(doc, "5 การอ่านแจ้งเตือนและรายละเอียดคำร้อง", "ระบบแจ้งความเคลื่อนไหวเมื่อมีการรับเรื่อง มอบหมายช่าง เปลี่ยนสถานะ หรือปิดงาน จุดสีแดงหมายถึงมีรายการที่ยังไม่ได้อ่าน", ["เปิดเมนู แจ้งเตือน", "เลือกการแจ้งเตือนเพื่อเปิดคำร้องที่เกี่ยวข้อง", "ตรวจสถานะ ผู้รับผิดชอบ ประวัติการดำเนินงาน และภาพก่อนหรือหลังซ่อม", "เมื่องานเสร็จ ผู้แจ้งตรวจผลงานและให้คะแนนพร้อมความคิดเห็นได้"], "05-reporter-notifications.png", "ภาพที่ 5 ศูนย์แจ้งเตือนของระบบ")
section(doc, "6 รายละเอียดคำร้องสำหรับผู้แจ้ง", "หน้ารายละเอียดรวบรวมข้อมูลทั้งหมดของคำร้องไว้ในจุดเดียว และแสดงขั้นตอนล่าสุดด้วยข้อความที่เข้าใจง่าย", ["ตรวจรหัสคำร้องและสถานะปัจจุบัน", "ตรวจสถานที่ ข้อมูลผู้แจ้ง รายละเอียดปัญหา และรูปก่อนซ่อม", "ดูชื่อช่างและประวัติการเปลี่ยนสถานะ", "หลังสถานะเสร็จสิ้น ให้ตรวจภาพหลังซ่อมก่อนประเมินบริการ"], "06-reporter-detail.png", "ภาพที่ 6 หน้ารายละเอียดคำร้องสำหรับผู้แจ้ง")
section(doc, "7 Dashboard และรายงานสำหรับผู้ดูแลระบบ", "Dashboard ใช้ติดตามจำนวนคำร้อง ระยะเวลาดำเนินการ ผลการให้บริการ และภาระงานของช่าง", ["เลือกช่วงวันที่ สถานะ อาคาร หมวดหมู่ หรือช่างที่ต้องการ", "กด แสดงผล เพื่อประมวลผลตามตัวกรอง", "ตรวจการ์ดสรุปและรายการสถิติ", "เลือกดาวน์โหลด CSV หรือพิมพ์ PDF เมื่อต้องนำข้อมูลไปใช้ต่อ"], "07-admin-dashboard.png", "ภาพที่ 7 Dashboard และรายงานสำหรับผู้ดูแลระบบ")
section(doc, "8 ศูนย์ควบคุมและรายการคำร้อง", "หน้าจัดการคำร้องประกอบด้วยแผนที่ ตัวกรอง และการ์ดคำร้อง ผู้ดูแลใช้ค้นหาและเปิดงานได้โดยไม่ต้องเลื่อนตารางแนวนอน", ["ตรวจหมุดบนแผนที่เพื่อดูจุดที่เกิดปัญหา", "ค้นหาด้วยรหัส หัวข้อ ผู้แจ้ง หรือสถานที่", "กรองตามหมวดหมู่และสถานะ", "เลือก ดูรายละเอียด หรือปุ่มจ่ายงานช่างจากการ์ดคำร้อง"], "08-admin-tickets.png", "ภาพที่ 8 ศูนย์ควบคุมและรายการคำร้องสำหรับผู้ดูแลระบบ")
section(doc, "9 การรับเรื่องและมอบหมายช่าง", "ในหน้ารายละเอียด ผู้ดูแลสามารถรับเรื่องพร้อมมอบหมายช่าง หรือเปลี่ยนช่างสำหรับงานที่กำลังดำเนินการได้โดยตรง", ["เปิดคำร้องที่ต้องการ", "เลือกช่างผู้รับผิดชอบและระดับความเร่งด่วน", "กรอกหมายเหตุถึงช่างหากมีเงื่อนไขหน้างาน", "กด รับเรื่องและมอบหมายช่าง สำหรับคำร้องใหม่ หรือยืนยันการเปลี่ยนช่างสำหรับงานเดิม", "หากไม่ดำเนินการให้เลือก ไม่รับดำเนินการ พร้อมระบุเหตุผล หรือเลือก ยกเลิกคำร้อง ตามกรณี"], "09-admin-detail.png", "ภาพที่ 9 การรับเรื่อง มอบหมายช่าง และจัดการสถานะคำร้อง")
section(doc, "10 การจัดการข้อมูลอาคารและหมวดหมู่", "ผู้ดูแลใช้หน้าข้อมูลพื้นฐานเพื่อเพิ่มหรือแก้ไขอาคาร พิกัดเริ่มต้น และหมวดหมู่ปัญหาที่ใช้ในแบบฟอร์ม", ["เปิดเมนู ข้อมูลพื้นฐาน หรือหน้าตั้งค่า", "เพิ่มหรือแก้ไขชื่ออาคาร รหัส และพิกัด", "เพิ่มหรือแก้ไขชื่อหมวดหมู่พร้อมคำอธิบาย", "บันทึกและตรวจว่าข้อมูลใหม่ปรากฏในแบบฟอร์มสร้างคำร้อง"], "10-admin-settings.png", "ภาพที่ 10 หน้าจัดการข้อมูลอาคารและหมวดหมู่")
section(doc, "11 รายการงานของช่าง", "ช่างเห็นเฉพาะงานที่ได้รับมอบหมายให้บัญชีของตน การ์ดงานแสดงสถานะ ความเร่งด่วน สถานที่ และข้อมูลที่จำเป็นก่อนลงพื้นที่", ["เปิดเมนู งานของฉัน", "ตรวจสถานะและความเร่งด่วนของแต่ละงาน", "เลือกงานเพื่อเปิดรายละเอียด", "อ่านรายละเอียดปัญหา สถานที่ รูปก่อนซ่อม และข้อมูลผู้เกี่ยวข้อง"], "11-technician-jobs.png", "ภาพที่ 11 รายการงานที่ได้รับมอบหมายสำหรับช่าง")
section(doc, "12 การอัปเดตงานซ่อมของช่าง", "ช่างที่ได้รับมอบหมายสามารถบันทึกความคืบหน้า เปลี่ยนสถานะ และแนบภาพหลังซ่อมได้", ["เลือกสถานะ กำลังดำเนินการ รออะไหล่ หรือซ่อมเสร็จ", "กรอกผลการดำเนินงานหรือเหตุผลที่รออะไหล่", "เมื่อซ่อมเสร็จ ให้แนบภาพหลังซ่อมที่เห็นผลชัดเจน", "กด บันทึกและอัปเดตสถานะงาน แล้วตรวจข้อความยืนยัน"], "12-technician-detail.png", "ภาพที่ 12 แผงบันทึกงานและอัปเดตสถานะสำหรับช่าง")
section(doc, "13 การใช้งานบนโทรศัพท์", "ระบบรองรับหน้าจอขนาดเล็ก เมนูหลักอยู่ด้านล่างและปุ่มดำเนินการมีขนาดเหมาะกับการสัมผัส", ["ใช้เมนูด้านล่างเพื่อสลับระหว่างหน้าหลัก แจ้งปัญหา คำร้อง และแจ้งเตือน", "กรอกแบบฟอร์มทีละขั้นตอนและเลื่อนหน้าลงตามลำดับ", "ปุ่ม ถัดไป อยู่ท้ายเนื้อหาของแต่ละขั้นตอน", "ตรวจข้อมูลและรูปภาพก่อนส่ง โดยหลีกเลี่ยงการปิดหน้าในระหว่างอัปโหลด"], "13-mobile-report.png", "ภาพที่ 13 แบบฟอร์มแจ้งปัญหาบนโทรศัพท์", width=3.0)

page(doc)
heading(doc, "14 ความหมายของสถานะคำร้อง", 2)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
for i, w in enumerate([1.35, 2.15, 2.7]): table.columns[i].width = Inches(w)
headers = ["สถานะ", "ความหมาย", "ผู้ดำเนินการถัดไป"]
for i, text in enumerate(headers): shade(table.rows[0].cells[i], "1E293B"); set_cell(table.rows[0].cells[i], text, True, True)
rows = [
    ("รอรับเรื่อง", "ส่งคำร้องแล้วและรอเจ้าหน้าที่ตรวจสอบ", "Admin รับเรื่องหรือปฏิเสธ"),
    ("กำลังดำเนินการ", "มอบหมายช่างแล้วและกำลังตรวจสอบหรือซ่อม", "ช่างอัปเดตความคืบหน้า"),
    ("รออะไหล่", "หยุดรออะไหล่ วัสดุ หรือการจัดซื้อ", "ช่างกลับมาอัปเดตเมื่อพร้อม"),
    ("เสร็จสิ้น", "ช่างบันทึกผลและปิดงานแล้ว", "ผู้แจ้งตรวจและประเมิน"),
    ("ไม่รับดำเนินการ", "เจ้าหน้าที่พิจารณาแล้วไม่สามารถดำเนินการ", "ตรวจเหตุผลหรือเปิดงานใหม่"),
    ("ยกเลิกแล้ว", "ผู้แจ้งหรือผู้ดูแลยกเลิกคำร้อง", "Admin เปิดงานใหม่ได้หากจำเป็น"),
]
for n, row in enumerate(rows):
    cells = table.add_row().cells
    if n % 2: [shade(c, "F1F5F9") for c in cells]
    for i, text in enumerate(row): set_cell(cells[i], text, i == 0)

heading(doc, "15 แนวทางแก้ปัญหาเบื้องต้น", 2)
items = [
    ("เข้าสู่ระบบไม่ได้", "ตรวจอีเมลและรหัสผ่าน ปิด Caps Lock แล้วลองใหม่ หากยังไม่ได้ให้ติดต่อผู้ดูแลระบบ"),
    ("แผนที่ไม่แสดง", "ตรวจอินเทอร์เน็ต รีเฟรชหน้า และอนุญาตตำแหน่ง หากยังไม่พบพิกัดให้กรอกสถานที่เป็นข้อความ"),
    ("แนบรูปไม่ได้", "ใช้ไฟล์ JPG PNG หรือ WebP ขนาดไม่เกิน 5 MB และลองเลือกรูปใหม่"),
    ("ข้อมูลเปลี่ยนแปลงแล้ว", "รีเฟรชหน้าเพื่อโหลดสถานะล่าสุดก่อนกดดำเนินการอีกครั้ง"),
    ("ไม่พบปุ่มที่ต้องการ", "ตรวจบทบาทบัญชีและสถานะคำร้อง เพราะปุ่มจะแสดงตามสิทธิ์และขั้นตอนงาน"),
    ("หน้าเว็บแสดงเวอร์ชันเก่า", "กด Ctrl Shift R บนคอมพิวเตอร์ หรือปิดแล้วเปิดหน้าใหม่บนโทรศัพท์"),
]
for title, detail in items:
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    font(p.add_run(f"• {title}: "), 15, True); font(p.add_run(detail), 15)

paragraph(doc, "เมื่อพบข้อผิดพลาดที่แก้ไม่ได้ ให้บันทึกรหัสคำร้อง หน้าที่พบปัญหา เวลาที่เกิดเหตุ และภาพหน้าจอ แล้วส่งข้อมูลให้ผู้ดูแลระบบเพื่อช่วยตรวจสอบ", 15, False, WD_ALIGN_PARAGRAPH.JUSTIFY, 6, 0)

doc.save(OUTPUT)
print(OUTPUT)
